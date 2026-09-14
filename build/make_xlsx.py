# -*- coding: utf-8 -*-
import sys, os, re, json, glob, csv, collections, statistics
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import d_config as C, d_arch as AR
import d_art_01,d_art_02,d_art_03,d_art_04,d_art_05,d_art_06,d_art_07
import d_blocks_01,d_blocks_02,d_blocks_03,d_blocks_04,d_blocks_05,d_blocks_06
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ARTS=[]
for m in (d_art_01,d_art_02,d_art_03,d_art_04,d_art_05,d_art_06,d_art_07): ARTS+=m.ART
BLOCKS={}
for m in (d_blocks_01,d_blocks_02,d_blocks_03,d_blocks_04,d_blocks_05,d_blocks_06): BLOCKS.update(m.B)
S=json.load(open(os.path.join(ROOT,'build/semantics.json'),encoding='utf-8'))
SEM=S['SEM']
PAGES=json.load(open(os.path.join(ROOT,'raw/pages_analysis.json'),encoding='utf-8'))
HUB={h[0]:h for h in AR.HUBS}
ART={a[0]:a for a in ARTS}
SEM_BY_PAGE=collections.defaultdict(list)
for r in SEM:
    if r['status']=='принято': SEM_BY_PAGE[r['page']].append(r)

wb=Workbook(); wb.remove(wb.active)
HDR=Font(bold=True,color='FFFFFF',size=10); HFILL=PatternFill('solid',fgColor='2F5597')
WRAP=Alignment(vertical='top',wrap_text=True); TOP=Alignment(vertical='top')
THIN=Border(*[Side(style='thin',color='D9D9D9')]*4)

def sheet(name, headers, rows, widths=None, freeze='A2'):
    ws=wb.create_sheet(name[:31])
    ws.append(headers)
    for c in ws[1]:
        c.font=HDR; c.fill=HFILL; c.alignment=Alignment(vertical='center',wrap_text=True)
    for r in rows: ws.append(list(r))
    ws.freeze_panes=freeze
    ws.auto_filter.ref=f'A1:{get_column_letter(len(headers))}{max(ws.max_row,1)}'
    ws.row_dimensions[1].height=32
    for i,h in enumerate(headers,1):
        w=(widths[i-1] if widths and i-1<len(widths) else 22)
        ws.column_dimensions[get_column_letter(i)].width=w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment=WRAP; c.border=THIN
    return ws

# ===== 02 Паспорт =====
sheet('02 Паспорт',['Параметр','Значение','Источник / владелец','Статус','Допущение','Дата'],
      C.PASSPORT,[30,70,34,16,46,12])

# ===== 28 Источники =====
sheet('28 Источники и сбор',['source_id','Источник','Файл / URL','Дата','Период данных','Настройки','Маркер / охват','Пакет','Сохранённые строки','Ошибки','Ограничения','Вклад новых фраз'],
      C.SOURCES,[10,34,26,16,24,34,18,16,26,34,54,24])

# ===== 09 Семантика =====
sheet('09 Семантика',['query_id','Исходная фраза','Нормализованная фраза','Широкая частотность, МО','Интент','Геозависимость по формулировке','Статус','Основание исключения','primary_page_id','Роль запроса на странице','source_id'],
      [(r['query_id'],r['raw'],r['norm'],r['freq'],r['intent'],r['geo'],r['status'],r['reason'],r['page'],r['role'],r['src']) for r in SEM],
      [10,44,44,14,20,14,12,44,14,40,18])

# ===== 11 Частотность =====
NARROW={}
np_=os.path.join(ROOT,'raw/wordstat/narrow.tsv')
if os.path.exists(np_):
    for i,l in enumerate(open(np_,encoding='utf-8')):
        if i==0 or '\t' not in l: continue
        a,b=l.rstrip('\n').rsplit('\t',1)
        try: NARROW[re.sub(r'\s+',' ',a.lower().replace('ё','е')).strip()]=int(b)
        except: pass
freq_rows=[]
for r in SEM:
    freq_rows.append((f"M{r['query_id'][1:]}", r['query_id'], r['raw'], 'широкая (без операторов)', r['freq'],
                      'Москва и область', C.PERIOD, 'все (десктоп+смартфон+планшет)', C.DATE, 'S01', 'измерено'))
nb=0
for r in SEM:
    key=r['norm']
    if key in NARROW:
        nb+=1
        v=NARROW[key]
        ratio=f"{r['freq']/v:.1f}x" if v else 'широкая/узкая не считается: узкая = 0'
        freq_rows.append((f"N{r['query_id'][1:]}", r['query_id'], '"'+r['raw']+'"', 'фразовая, в кавычках', v,
                          'Москва и область', C.PERIOD, 'все (десктоп+смартфон+планшет)', C.DATE, 'S12',
                          f'измерено; отношение широкой к узкой: {ratio}'))
for r in SEM:
    if r['status']=='принято' and r['norm'] not in NARROW:
        freq_rows.append((f"N{r['query_id'][1:]}", r['query_id'], '"'+r['raw']+'"', 'фразовая, в кавычках', None,
                          'Москва и область', C.PERIOD, 'все (десктоп+смартфон+планшет)', '', 'S12',
                          'НЕ ИЗМЕРЕНО — вне выборки замеров узкой частотности (измерены 78 приоритетных фраз)'))
sheet('11 Частотность',['measurement_id','query_id','Строка измерения','Тип частотности','Значение','Регион','Период','Устройство','Дата','source_id','Статус'],
      freq_rows,[12,10,44,26,12,20,24,26,12,10,14])

# ===== 12 Сезонность =====
seas=[]
if os.path.exists(os.path.join(ROOT,'raw/seasonality/seasonality.tsv')):
    for row in csv.DictReader(open(os.path.join(ROOT,'raw/seasonality/seasonality.tsv'),encoding='utf-8'),delimiter='\t'):
        seas.append((row['query'],row['month'],int(row['value']),'число запросов, широкая','Москва и область','S02','24 мес., полный ряд',''))
bym=collections.defaultdict(dict)
for q,m,v,*_ in seas: bym[q][m]=v
for q,d in bym.items():
    lo=min(d,key=d.get); hi=max(d,key=d.get)
    for row in seas:
        if row[0]==q and row[1]==hi:
            seas[seas.index(row)]=row[:7]+(f'ПИК: {hi}={d[hi]}; ДНО: {lo}={d[lo]}; размах {d[hi]/max(d[lo],1):.1f}x',)
sheet('12 Сезонность',['query_id / запрос','Месяц','Значение','Тип измерения','Регион','source_id','Полнота ряда','Вывод'],
      seas,[34,12,12,26,20,10,22,60])

# ===== 03 Структура =====
struct=[]
for c in AR.COMMERCIAL:
    struct.append((c[0],c[1],c[2],c[3],c[4],c[5],c[6],c[7],c[8],c[9],c[10],c[11]))
for h in AR.HUBS:
    struct.append((h[0],h[1],'P506','инфо-хаб',h[0],h[2],h[4],h[5],f"Опорный запрос «{h[6]}» — {h[7]} показов/мес по МО; в хабе {h[8]} статей",'новый','1','создать'))
for a in ARTS:
    h=HUB[a[1]]
    struct.append((a[0],h[1]+a[2]+'/',a[1],'статья',a[1],a[3],a[4],a[5],a[9],'новый',str(a[11]),'ТЗ подготовлено'))
sheet('03 Структура',['page_id','URL','parent_page_id','Тип','Кластер','H1','Title','Description','Основание','Существующий / новый','Очередь','Статус'],
      struct,[10,46,14,14,10,46,56,72,56,20,10,18])

# ===== 04 План статей =====
plan=[]
for a in ARTS:
    h=HUB[a[1]]
    cl=SEM_BY_PAGE.get(a[0],[])
    note = '' if cl else 'НЕТ ЗАКРЕПЛЁННЫХ ИЗМЕРЕННЫХ ФРАЗ: тема опирается на интент и выдачу, а не на замеренную частотность. Перед постановкой в работу проверить спрос дополнительным маркером'
    plan.append((a[0],a[3],h[2],'см. «Аудитория и задача»',a[9],a[10],a[19] if len(a)>19 else a[17],
                 a[11], f'{a[12]}–{a[13]} зн. с пробелами', len(cl), sum(x['freq'] for x in cl), a[6], a[7], note))
sheet('04 План статей',['page_id','Тема (H1)','Хаб','Аудитория','Задача пользователя','Отличие от соседних материалов','Связанная услуга / CTA','Приоритет','Трудоёмкость (объём)','Фраз в кластере','Сумма широких частотностей (с пересечениями)','Основной запрос','Частотность основного','Предупреждение'],
      plan,[10,46,26,26,56,66,34,10,24,12,20,34,14,80])

# ===== 05 ТЗ страниц =====
tz=[]
for a in ARTS:
    h=HUB[a[1]]
    cl=SEM_BY_PAGE.get(a[0],[])
    qids=', '.join(x['query_id'] for x in cl) or '—'
    tz.append((a[0],a[9],f'Раскрыть: {a[3]}. Оставить соседним материалам: смежные темы хаба «{h[2]}» (см. столбец «Входящие/исходящие ссылки»).',
               a[6],qids,f'Title: {a[4]} ({len(a[4])} зн.) | Description: {a[5]} ({len(a[5])} зн.) | H1: {a[3]}',
               f'{a[12]}–{a[13]} зн. с пробелами',
               'Диапазон назначен после составления плана H2–H3 и сопоставлен с измеренными объёмами сопоставимых страниц конкурентов (лист «18 Объёмы и рекомендации»)',
               a[14],a[15],a[16],a[17],'Автор: не назначен; редактор: не назначен; эксперт: владелец / инструктор базы (контакты не переданы)',
               'Приёмка: все H2 раскрыты; каждый факт из столбца «Факты на проверку» подтверждён или снят; нет придуманных цен и характеристик; внутренние ссылки проставлены; alt у всех изображений; объём в диапазоне',
               'ТЗ подготовлено'))
sheet('05 ТЗ страниц',['page_id','Задача пользователя','Границы темы','Основной запрос','Все query_id кластера','Метаданные (Title / Description / H1)','Диапазон объёма','Обоснование объёма','Собственная польза','Медиа и материалы','Факты на проверку','CTA','Автор / редактор / эксперт','Критерии приёмки','Статус'],
      tz,[10,56,70,34,70,90,24,70,60,50,70,40,50,90,16])

# ===== 06 План H2-H3 =====
bl=[]
for a in ARTS:
    pid=a[0]; parent=None; order=0
    for raw in BLOCKS[pid]:
        lvl,head,thesis = (raw.split('|',2)+['',''])[:3]
        order+=1
        bid=f'{pid}-B{order:02d}'
        if lvl=='2': parent=bid; par=''
        else: par=parent or ''
        bl.append((bid,pid,par,order,'H'+lvl,head,thesis,'',''))
sheet('06 План H2-H3',['block_id','page_id','parent_block_id','Порядок','Уровень','Готовый заголовок','Задача раздела и тезисы','Ключи (query_id)','Ориентир объёма'],
      bl,[16,10,16,9,9,60,90,20,16])

# ===== 07 Запросы в тексте =====
qt=[]
for r in SEM:
    if r['status']!='принято': continue
    pid=r['page']
    b=f'{pid}-B01'
    if r['role']=='основной': way='отдельный ответ в первом экране и в блоке «Коротко о главном»'
    elif r['role']=='поддерживающий': way='часть профильного раздела H2'
    else: way='покрыт синонимичной формулировкой внутри раздела; дословное вхождение не требуется'
    lit='да, в Title/H1' if r['role']=='основной' else 'нет — естественная формулировка'
    qt.append((r['query_id'],pid,b,r['intent'],way,r['raw'],lit))
sheet('07 Запросы в тексте',['query_id','page_id','block_id','Задача запроса','Способ покрытия','Рекомендуемая естественная формулировка','Дословное употребление и причина'],
      qt,[10,10,14,24,60,44,44])

# ===== 10 Кластеры =====
cls=[]
for pid,rows in sorted(SEM_BY_PAGE.items(), key=lambda x:-len(x[1])):
    name = ART[pid][3] if pid in ART else next((c[5] for c in AR.COMMERCIAL if c[0]==pid), pid)
    main=max(rows,key=lambda x:x['freq'])
    checked = 'да' if pid in ('P003','P004','P005','P010','P017','A001','A106','A117','A130','P007','A035','A021','A036') else 'нет — выдача по опорному запросу кластера не снималась'
    cls.append((f'CL-{pid}',name,'см. «04 План статей»','инфо' if pid.startswith('A') else 'коммерческий',
                pid,main['query_id'],len(rows),checked,'Границы заданы основным запросом статьи и её H2-планом',
                'Объединение по пересечению лексики основного и поддерживающих запросов','—'))
sheet('10 Кластеры',['cluster_id','Название','Задача','Формат','page_id','Основной query_id','Число фраз','Полнота проверки выдачи','Границы','Основание объединения','Спорные случаи'],
      cls,[14,46,20,16,10,14,12,46,50,60,20])

# ===== 13 Выдача =====
serp=[]
sp=os.path.join(ROOT,'raw/serp/serp.tsv')
if os.path.exists(sp):
    n=0
    for row in csv.DictReader(open(sp,encoding='utf-8'),delimiter='\t'):
        n+=1
        serp.append((f'SR{n:04d}',row['query'],'13–14.09.2026','Москва (lr=1)','десктоп',int(row['pos']),row['url'],row['domain'],
                     'коммерческая' if not re.search(r'/blog|/stati|/article|news',row['url']) else 'информационная',
                     'органика','S03','полный топ-10 органики; реклама (yabs) и блок Карт исключены из нумерации'))
sheet('13 Выдача',['serp_id','query_id / запрос','Дата','Регион','Устройство','Позиция','URL','Домен','Тип страницы','Тип результата','source_id','Полнота'],
      serp,[10,44,16,16,12,9,70,26,18,14,10,60])

# ===== 15 Контент конкурентов =====
cont=[]
for p in PAGES:
    if p.get('status')!='ok':
        cont.append((p.get('content_id',''),p.get('url',''),C.DATE,'curl + собственный парсер','нет доступа / пустой ответ','','','','','','','','','','','','НЕ ИЗМЕРЕНО — ошибка доступа, нулевой объём не проставлялся'))
        continue
    cont.append((p['content_id'],p['url'],C.DATE,'curl + парсер: исключены script/style/nav/header/footer/form','ok',
                 p['words'],p['chars_sp'],p['chars_nosp'],p['title'],p['description'],'; '.join(p['h1']),
                 len(p['h2']),len(p['h3']),p['imgs'],p['tables'],p['videos'],
                 'FAQ-разметка: '+('есть' if p['faq_schema'] else 'нет')+'; цены в тексте: '+('есть' if p['has_price'] else 'нет')))
sheet('15 Контент конкурентов',['content_id','URL','Дата','Метод извлечения','Статус','Слова','Знаки с пробелами','Знаки без пробелов','Title','Description','H1','Кол-во H2','Кол-во H3','Изображения','Таблицы','Видео','Примечание'],
      cont,[12,66,12,54,16,9,14,14,60,72,50,10,10,12,10,9,50])

# ===== 16 Заголовки конкурентов =====
hd=[]
for p in PAGES:
    if p.get('status')!='ok': continue
    o=0; parent=''
    for h in p['h2']:
        o+=1; hid=f"{p['content_id']}-H{o:02d}"; parent=hid
        hd.append((p['content_id'],hid,'',o,'H2',h,'',''))
    for h in p['h3']:
        o+=1; hid=f"{p['content_id']}-H{o:02d}"
        hd.append((p['content_id'],hid,parent,o,'H3',h,'','порядок H3 восстановлен отдельным списком; точная вложенность в исходной разметке не всегда определима'))
sheet('16 Заголовки конкурентов',['content_id','heading_id','parent_heading_id','Порядок','Уровень','Текст заголовка','Тема','Примечание'],
      hd,[12,16,16,9,9,70,24,60])

# ===== 18 Объёмы и рекомендации =====
ok=[p for p in PAGES if p.get('status')=='ok']
def stat(sel, label, pid):
    if len(sel)<1: return None
    w=[p['words'] for p in sel]; ch=[p['chars_sp'] for p in sel]
    dom=len({re.sub(r'^www\.','',re.sub(r'^https?://','',p['url']).split('/')[0]) for p in sel})
    p25=p75=''
    if len(sel)>=5:
        p25=round(statistics.quantiles(ch,n=4)[0]); p75=round(statistics.quantiles(ch,n=4)[2])
    return (pid,label,len(sel),dom,min(w),max(w),round(statistics.mean(w)),round(statistics.median(w)),
            min(ch),max(ch),round(statistics.mean(ch)),round(statistics.median(ch)),p25,p75,
            'слова и знаки с пробелами',
            (f'{p25}–{p75} зн. с пробелами (P25–P75 наблюдаемых объёмов). Верхняя граница — для материалов с планом от 12 блоков H2' if p25
             else f'{min(ch)}–{max(ch)} зн. с пробелами (весь наблюдаемый разброс; квартили не считались из-за малого n)'),
            'ограниченная выборка (n<3)' if len(sel)<3 else ('P25–P75 по методу quantiles(n=4)' if len(sel)>=5 else 'без квартилей: n<5'))
groups=[('Коммерческие страницы проката квадроциклов', lambda u: re.search(r'kvadro|atv|quad',u) and not re.search(r'/blog|corporate|korporativ',u), 'P003'),
        ('Корпоративные и тимбилдинговые страницы', lambda u: re.search(r'corporate|korporativ|timbilding|team',u), 'P010'),
        ('Информационные статьи блога', lambda u: re.search(r'/blog',u), 'A001'),
        ('Страницы снегоходов', lambda u: re.search(r'snow|snegohod',u), 'P004'),
        ('Страницы багги', lambda u: re.search(r'buggy|baggi',u), 'P005'),
        ('Детские и семейные страницы', lambda u: re.search(r'detsk|family|children',u), 'P007')]
vol=[]
for label,f,pid in groups:
    sel=[p for p in ok if f(p['url'])]
    s=stat(sel,label,pid)
    if s: vol.append(s)
sheet('18 Объёмы и рекомендации',['page_id (пример применения)','Выборка','n','Независимых доменов','min слов','max слов','среднее слов','медиана слов','min знаков','max знаков','среднее знаков','медиана знаков','P25 знаков','P75 знаков','Единицы','Рекомендуемый диапазон','Ограничения'],
      vol,[22,50,6,10,10,10,12,12,11,11,12,13,11,11,22,26,46])

wb.save(os.path.join(ROOT,'out/_partial.xlsx'))
print('часть 1 сохранена, листов:',len(wb.sheetnames))
print(wb.sheetnames)
