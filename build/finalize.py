# -*- coding: utf-8 -*-
import os
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill
ROOT=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
f=os.path.join(ROOT,'out','SEO-прокат-внедорожной-техники-Москва-и-область-2026-09-14.xlsx')
wb=load_workbook(f)
def dv(sheet, col, values, n):
    ws=wb[sheet]
    d=DataValidation(type='list', formula1='"'+','.join(values)+'"', allow_blank=True, showDropDown=False)
    ws.add_data_validation(d); d.add(f'{col}2:{col}{n}')
dv('09 Семантика','G',['принято','исключено','резерв','требует проверки'],wb['09 Семантика'].max_row)
dv('03 Структура','L',['создать','обновить','объединить','убрать из индекса','опубликовано'],wb['03 Структура'].max_row)
dv('05 ТЗ страниц','O',['ТЗ подготовлено','ТЗ принято','материалы запрошены','текст','редактура','экспертная проверка','загрузка','опубликовано','проверено после запуска'],wb['05 ТЗ страниц'].max_row)
dv('25 Производство','I',['ТЗ подготовлено','ТЗ принято','материалы запрошены','текст','редактура','экспертная проверка','загрузка','опубликовано','проверено после запуска'],wb['25 Производство'].max_row)
dv('22 Техника и CMS','J',['к работе','в работе','выполнено','отложено','не применимо'],wb['22 Техника и CMS'].max_row)
dv('19 Материалы и факты','I',['запрошено','получено','проверено','отклонено','БЛОКИРУЕТ ПУБЛИКАЦИЮ'],wb['19 Материалы и факты'].max_row)
# подсветка блокирующих строк
red=PatternFill('solid',fgColor='FFE0E0')
for sh,col in (('01 Резюме',4),('32 Приёмка',5),('19 Материалы и факты',9)):
    ws=wb[sh]
    for r in ws.iter_rows(min_row=2):
        v=str(r[col-1].value or '')
        if any(k in v for k in ('НЕ ВЫПОЛНЕН','ЗАБЛОКИР','БЛОКИРУЕТ','НЕ ПРОВЕРЕНО','РЕЖИМ ИЗМЕНЁН','ДА —','НЕ ДОСТИГНУТ','РАСХОЖДЕНИЕ')):
            for c in r: c.fill=red
wb.save(f); print('финализировано:',f)
